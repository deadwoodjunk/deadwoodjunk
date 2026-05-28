"""営業に必要な3つの知識 ─ PowerPoint(.pptx) 生成スクリプト

HTMLスライド(sales_knowledge_framework.html)と同じ構成・配色で
16:9のスライドを生成する。
    python3 build_slides.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- カラーパレット（HTML版と統一） ----
C1 = RGBColor(0x25, 0x63, 0xEB)  # 業界知識（青）
C2 = RGBColor(0x16, 0xA3, 0x4A)  # 顧客個別の情報（緑）
C3 = RGBColor(0xEA, 0x58, 0x0C)  # 自社製品知識（オレンジ）
INK = RGBColor(0x1E, 0x29, 0x3B)
SUB = RGBColor(0x64, 0x74, 0x8B)
BG = RGBColor(0xF1, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x0F, 0x17, 0x2A)
LINE = RGBColor(0xE2, 0xE8, 0xF0)

FONT = "Meiryo"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    rect = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    rect.fill.solid()
    rect.fill.fore_color.rgb = bg
    rect.line.fill.background()
    rect.shadow.inherit = False
    return s


def txt(slide, left, top, width, height, text, size, color=INK, bold=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT, spacing=None,
        line_spacing=None):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
        if spacing is not None:
            _set_spacing(r, spacing)
    return tb


def _set_spacing(run, pts):
    """文字間隔(トラッキング)を設定。"""
    rPr = run._r.get_or_add_rPr()
    rPr.set("spc", str(int(pts * 100)))


def rounded(slide, left, top, width, height, fill, line=None, radius=0.08):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    return sh


def circle(slide, left, top, size, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def set_alpha(shape, alpha_pct):
    """図形の塗りに透明度を設定（ベン図の重なり表現用）。"""
    sp = shape.fill._xPr.find(qn('a:solidFill'))
    srgb = sp.find(qn('a:srgbClr'))
    a = srgb.makeelement(qn('a:alpha'), {'val': str(int(alpha_pct * 1000))})
    srgb.append(a)


# ============================================================
# Slide 1 ─ タイトル
# ============================================================
s = add_slide()
txt(s, Inches(1), Inches(1.5), Inches(11.33), Inches(0.5),
    "SALES KNOWLEDGE FRAMEWORK", 14, SUB, bold=True,
    align=PP_ALIGN.CENTER, spacing=3)
txt(s, Inches(1), Inches(2.1), Inches(11.33), Inches(2),
    "営業に必要な3つの知識", 54, INK, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(2), Inches(3.9), Inches(9.33), Inches(1),
    "成果を出す営業は「何を売るか」だけでなく、顧客と市場を立体的に理解している。\nその土台となる3要素を整理します。",
    16, SUB, align=PP_ALIGN.CENTER, line_spacing=1.4)

pills = [("1", "業界知識", C1), ("2", "顧客個別の情報", C2), ("3", "自社製品知識", C3)]
pw, gap = Inches(3.3), Inches(0.35)
total = pw * 3 + gap * 2
x = (SW - total) // 2
py = Inches(5.4)
for num, label, col in pills:
    pill = rounded(s, x, py, pw, Inches(0.85), WHITE, radius=0.5)
    circle(s, x + Inches(0.25), py + Inches(0.2), Inches(0.45), col)
    txt(s, x + Inches(0.25), py + Inches(0.2), Inches(0.45), Inches(0.45),
        num, 16, WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x + Inches(0.8), py, pw - Inches(0.9), Inches(0.85),
        label, 18, INK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    x += pw + gap


# ============================================================
# Slide 2 ─ 全体像（ベン図）
# ============================================================
s = add_slide()
txt(s, Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4),
    "OVERVIEW ─ 全体像", 13, SUB, bold=True, spacing=2)
txt(s, Inches(0.8), Inches(0.95), Inches(11.7), Inches(0.7),
    "3つが重なる中心に「刺さる提案」が生まれる", 26, INK, bold=True)
txt(s, Inches(0.8), Inches(1.7), Inches(11.7), Inches(0.9),
    "どれか1つが欠けても提案は弱くなる。3つの知識が交わる領域こそが、顧客に選ばれる「最適提案」のスイートスポット。",
    14, SUB, line_spacing=1.3)

# ベン図（3円・透過で重ね合わせ）
cd = Inches(3.3)
cx = SW // 2
top_y = Inches(2.6)
top = circle(s, cx - cd // 2, top_y, cd, C1)
left = circle(s, cx - cd + Inches(0.55), top_y + Inches(1.5), cd, C2)
right = circle(s, cx - Inches(0.55), top_y + Inches(1.5), cd, C3)
for c in (top, left, right):
    set_alpha(c, 78)
# ラベル
txt(s, cx - cd // 2, top_y + Inches(0.45), cd, Inches(0.5),
    "① 業界知識", 16, WHITE, bold=True, align=PP_ALIGN.CENTER)
txt(s, cx - cd + Inches(0.55), top_y + Inches(2.7), cd, Inches(0.5),
    "② 顧客個別の情報", 15, WHITE, bold=True, align=PP_ALIGN.CENTER)
txt(s, cx - Inches(0.55), top_y + Inches(2.7), cd, Inches(0.5),
    "③ 自社製品知識", 15, WHITE, bold=True, align=PP_ALIGN.CENTER)
# 中心
txt(s, cx - Inches(0.9), top_y + Inches(1.85), Inches(1.8), Inches(0.7),
    "刺さる提案", 16, WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# ============================================================
# Slide 3 ─ 3要素サマリー（カード）
# ============================================================
s = add_slide()
txt(s, Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4),
    "3 ELEMENTS ─ 要素一覧", 13, SUB, bold=True, spacing=2)
txt(s, Inches(0.8), Inches(0.95), Inches(11.7), Inches(0.7),
    "3つの知識でカバーする範囲", 26, INK, bold=True)

cards = [
    ("1", "業界知識", "顧客が属する「市場」を読む", C1,
     ["市場トレンド・将来予測", "競合・業界構造", "法規制・ガイドライン", "業界特有の課題・専門用語"]),
    ("2", "顧客個別の情報", "目の前の「相手」を知る", C2,
     ["経営課題・現場の困りごと", "組織図・意思決定プロセス", "キーパーソン／決裁者", "予算・導入時期・購買履歴"]),
    ("3", "自社製品知識", "提供できる「価値」を語る", C3,
     ["機能・仕様・できること", "強み／弱み・差別化ポイント", "価格・導入条件・サポート", "導入事例・成功実績"]),
]
cw, cgap = Inches(3.85), Inches(0.4)
ctotal = cw * 3 + cgap * 2
cx0 = (SW - ctotal) // 2
cy = Inches(2.0)
ch = Inches(4.6)
for num, title, cap, col, items in cards:
    card = rounded(s, cx0, cy, cw, ch, WHITE, radius=0.05)
    # 上部カラーバー
    bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx0, cy, cw, Inches(0.12), )
    bar.fill.solid(); bar.fill.fore_color.rgb = col; bar.line.fill.background()
    bar.shadow.inherit = False
    # バッジ
    badge = rounded(s, cx0 + Inches(0.35), cy + Inches(0.45), Inches(0.55), Inches(0.55), col, radius=0.25)
    txt(s, cx0 + Inches(0.35), cy + Inches(0.45), Inches(0.55), Inches(0.55),
        num, 18, WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, cx0 + Inches(1.05), cy + Inches(0.45), cw - Inches(1.2), Inches(0.55),
        title, 19, INK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, cx0 + Inches(0.35), cy + Inches(1.15), cw - Inches(0.7), Inches(0.4),
        cap, 12, SUB)
    # 箇条書き
    iy = cy + Inches(1.7)
    for it in items:
        circle(s, cx0 + Inches(0.4), iy + Inches(0.12), Inches(0.12), col)
        txt(s, cx0 + Inches(0.7), iy, cw - Inches(1.0), Inches(0.5),
            it, 13.5, INK, anchor=MSO_ANCHOR.TOP)
        iy += Inches(0.68)
    cx0 += cw + cgap


# ============================================================
# Slide 4-6 ─ 各要素の詳細
# ============================================================
details = [
    ("01", "1", "業界知識", C1,
     ["市場規模・成長性とトレンド", "競合他社の動向と業界構造",
      "法規制・コンプライアンス", "業界用語・商習慣の理解"],
     "役割：顧客と「同じ目線」で語り、信頼を獲得する。話の前提が共有できる営業は強い。"),
    ("02", "2", "顧客個別の情報", C2,
     ["顧客固有の課題・ニーズ", "組織体制・意思決定フロー",
      "キーパーソンと関係構築", "予算規模・導入タイミング"],
     "役割：提案を「あなた向け」にカスタマイズする。一般論を“自分ごと”に変える鍵。"),
    ("03", "3", "自社製品知識", C3,
     ["機能・仕様を正確に説明できる", "強み・差別化ポイントの提示",
      "価格・契約・サポート条件", "導入事例で効果をイメージ化"],
     "役割：課題に対する「解決策」を具体的に示す。価値を翻訳して伝える力。"),
]
for eyebrow, num, title, col, items, role in details:
    s = add_slide()
    # 左：大きな番号ボックス
    box = rounded(s, Inches(1.0), Inches(2.0), Inches(3.5), Inches(3.5), col, radius=0.08)
    txt(s, Inches(1.0), Inches(2.0), Inches(3.5), Inches(3.5),
        num, 120, WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # 右：見出しと項目
    rx = Inches(5.2)
    rw = Inches(7.3)
    txt(s, rx, Inches(1.4), rw, Inches(0.4), "ELEMENT " + eyebrow, 13, SUB, bold=True, spacing=2)
    txt(s, rx, Inches(1.8), rw, Inches(0.8), title, 36, col, bold=True)
    iy = Inches(2.8)
    for it in items:
        txt(s, rx, iy, Inches(0.4), Inches(0.5), "✓", 16, col, bold=True)
        txt(s, rx + Inches(0.45), iy, rw - Inches(0.5), Inches(0.5), it, 16, INK, anchor=MSO_ANCHOR.MIDDLE)
        ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, rx, iy + Inches(0.52), rw, Pt(1))
        ln.fill.solid(); ln.fill.fore_color.rgb = LINE; ln.line.fill.background(); ln.shadow.inherit = False
        iy += Inches(0.62)
    rolebox = rounded(s, rx, iy + Inches(0.15), rw, Inches(0.9), WHITE, radius=0.12)
    txt(s, rx + Inches(0.3), iy + Inches(0.15), rw - Inches(0.6), Inches(0.9),
        "💡 " + role, 13.5, SUB, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.2)


# ============================================================
# Slide 7 ─ まとめ（掛け算）
# ============================================================
s = add_slide()
txt(s, Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.4),
    "SUMMARY ─ まとめ", 13, SUB, bold=True, spacing=2)
txt(s, Inches(0.8), Inches(1.15), Inches(11.7), Inches(0.8),
    "3つの掛け算が、提案の質を決める", 28, INK, bold=True)

# 数式: 業界 × 顧客 × 製品 ＝ 選ばれる提案力
fy = Inches(2.9)
fh = Inches(1.0)
elems = [("業界知識", C1, Inches(2.1)), ("顧客個別の情報", C2, Inches(2.5)),
         ("自社製品知識", C3, Inches(2.5))]
# レイアウトを計算
op_w = Inches(0.6)
result_w = Inches(2.9)
widths = [w for _, _, w in elems]
total_w = sum((w for w in widths), Emu(0)) + op_w * 3 + result_w
fx = (SW - total_w) // 2
for i, (label, col, w) in enumerate(elems):
    rounded(s, fx, fy, w, fh, col, radius=0.18)
    txt(s, fx, fy, w, fh, label, 17, WHITE, bold=True,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    fx += w
    op = "×" if i < len(elems) - 1 else "＝"
    txt(s, fx, fy, op_w, fh, op, 28, SUB, bold=True,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    fx += op_w
rounded(s, fx, fy, result_w, fh, DARK, radius=0.12)
txt(s, fx, fy, result_w, fh, "選ばれる提案力", 18, WHITE, bold=True,
    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

txt(s, Inches(1.5), Inches(4.5), Inches(10.33), Inches(1.5),
    "足し算ではなく「掛け算」。どれか1つでもゼロに近いと、提案全体の説得力が大きく下がる。\n"
    "3要素をバランスよく磨き続けることが、継続的に成果を出す営業の条件。",
    16, SUB, align=PP_ALIGN.CENTER, line_spacing=1.5)


prs.save("sales_knowledge_framework.pptx")
print("saved: sales_knowledge_framework.pptx  /  slides:", len(prs.slides._sldIdLst))
