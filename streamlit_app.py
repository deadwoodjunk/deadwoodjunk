"""Excelを投げ込んだら数式を自動設定してくれるStreamlitアプリ。"""
from __future__ import annotations

import hashlib
import io
import os

import anthropic
import streamlit as st
from openpyxl import load_workbook

from excel_agent import (
    apply_insertions,
    describe_workbook,
    generate_formula_plan,
)

st.set_page_config(page_title="Excel関数自動設定", page_icon="📊", layout="wide")

st.title("📊 Excel関数自動設定アプリ")
st.caption(
    "Excelをアップロードして自然言語で指示すると、Claudeが構造を解析して適切な数式を提案・挿入します。"
)

with st.sidebar:
    st.header("設定")
    api_key = st.text_input(
        "Anthropic API Key",
        type="password",
        value=os.environ.get("ANTHROPIC_API_KEY", ""),
        help="https://console.anthropic.com/ で取得できます。",
    )
    st.markdown("---")
    st.markdown("**モデル**: `claude-opus-4-7`")
    st.markdown("**思考モード**: adaptive thinking")
    st.markdown("**出力形式**: structured JSON")

uploaded = st.file_uploader("Excelファイル (.xlsx) をアップロード", type=["xlsx"])

if uploaded is not None:
    file_bytes = uploaded.getvalue()
    file_hash = hashlib.md5(file_bytes).hexdigest()

    if st.session_state.get("file_hash") != file_hash:
        st.session_state["file_hash"] = file_hash
        st.session_state["file_bytes"] = file_bytes
        st.session_state.pop("plan", None)
        st.session_state.pop("usage", None)

    try:
        wb_preview = load_workbook(io.BytesIO(file_bytes), data_only=False)
    except Exception as e:
        st.error(f"Excelの読み込みに失敗しました: {e}")
        st.stop()

    description = describe_workbook(wb_preview)

    with st.expander("📋 検出されたワークブック構造", expanded=False):
        st.code(description, language="text")

    instruction = st.text_area(
        "指示（自然言語）",
        placeholder=(
            "例: B列の売上を合計して最終行に出して、C列との比率をD列に入れて。\n"
            "空欄なら自動で実用的な集計を提案します。"
        ),
        height=120,
        key="instruction",
    )

    col_run, _ = st.columns([1, 4])
    with col_run:
        run = st.button(
            "🤖 数式を生成",
            type="primary",
            disabled=not api_key,
            use_container_width=True,
        )
    if not api_key:
        st.warning("サイドバーで Anthropic API Key を設定してください。")

    if run and api_key:
        with st.spinner("Claudeが数式を考えています..."):
            client = anthropic.Anthropic(api_key=api_key)
            try:
                plan, usage = generate_formula_plan(client, description, instruction)
            except anthropic.APIStatusError as e:
                st.error(f"API エラー ({e.status_code}): {e.message}")
                st.stop()
            except Exception as e:
                st.error(f"予期せぬエラー: {e}")
                st.stop()

        st.session_state["plan"] = plan
        st.session_state["usage"] = usage

if "plan" in st.session_state:
    plan = st.session_state["plan"]
    usage = st.session_state["usage"]

    st.markdown("---")
    st.subheader("🧠 提案された数式")

    summary = plan.get("summary", "").strip()
    if summary:
        st.info(summary)

    insertions = plan.get("insertions", [])
    if insertions:
        st.dataframe(
            insertions,
            use_container_width=True,
            hide_index=True,
            column_config={
                "sheet": st.column_config.TextColumn("シート", width="small"),
                "cell": st.column_config.TextColumn("セル", width="small"),
                "formula": st.column_config.TextColumn("数式", width="large"),
                "label": st.column_config.TextColumn("説明", width="medium"),
            },
        )
    else:
        st.warning("提案された数式がありません。指示を変えて再試行してください。")

    col_apply, _ = st.columns([1, 4])
    with col_apply:
        apply_btn = st.button(
            "✅ 適用してダウンロード",
            type="primary",
            disabled=not insertions,
            use_container_width=True,
        )

    with st.expander("使用トークン", expanded=False):
        st.json(
            {
                "input_tokens": usage.input_tokens,
                "output_tokens": usage.output_tokens,
                "cache_creation_input_tokens": getattr(usage, "cache_creation_input_tokens", 0),
                "cache_read_input_tokens": getattr(usage, "cache_read_input_tokens", 0),
            }
        )

    if apply_btn and insertions:
        wb = load_workbook(io.BytesIO(st.session_state["file_bytes"]), data_only=False)
        warnings = apply_insertions(wb, insertions)
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)

        for w in warnings:
            st.warning(w)

        st.success(f"{len(insertions) - len(warnings)} 個の数式を適用しました。")
        st.download_button(
            "📥 結果をダウンロード",
            data=out.getvalue(),
            file_name="output_with_formulas.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
        )
else:
    if uploaded is None:
        st.info("👆 上のフォームから Excel ファイルをアップロードしてください。")
