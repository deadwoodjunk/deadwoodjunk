"""Excel関数の自動設定ロジック。Claude APIに数式提案を依頼し、openpyxlで適用する。"""
from __future__ import annotations

import json
import re
from typing import Any

import anthropic
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

MODEL = "claude-opus-4-7"

SYSTEM_PROMPT = """あなたはExcelの数式設計エキスパートです。
ユーザーがアップロードしたワークブックの構造（シート名、ヘッダ、サンプル行、データ範囲）と
自然言語の指示を受け取り、適切なExcel数式をどのセルに挿入するかをJSONで返します。

# ルール
- Excelの組み込み関数を使用してください。例:
  SUM, AVERAGE, COUNT, COUNTA, COUNTIF, COUNTIFS, SUMIF, SUMIFS, IF, IFS, IFERROR,
  VLOOKUP, XLOOKUP, HLOOKUP, INDEX, MATCH, MAX, MIN, MEDIAN, LARGE, SMALL,
  ROUND, ROUNDUP, ROUNDDOWN, INT, MOD, ABS, SQRT, POWER,
  CONCAT, TEXTJOIN, LEN, LEFT, RIGHT, MID, FIND, SEARCH, SUBSTITUTE, TRIM, UPPER, LOWER,
  TODAY, NOW, DATE, YEAR, MONTH, DAY, WEEKDAY, EOMONTH, DATEDIF,
  AND, OR, NOT, ISBLANK, ISNUMBER, ISERROR, RANK, AVERAGEIF
- セル参照はA1形式。シート間参照は 'シート名'!A1 形式。シート名に空白や日本語があれば必ずシングルクォートで囲む。
- すべての数式は "=" で始める。
- 既存データを上書きしない位置（最終行+1、最右列+1、明らかな空きセル）を選ぶ。
- ヘッダ行（通常1行目）の真下や、データブロックの直下/右側に集計を置くのが自然。
- 各挿入には、なぜそこにその数式かを短い日本語の `label` で説明する。
- ユーザー指示が曖昧/未指定なら、データから推測される実用的な集計関数（合計、平均、件数、最大、最小、
  条件付き集計など）を5〜15個程度自由に提案する。
- 指示が具体的なら、それを最優先で実装し、関連する追加提案も少量加える。
"""

OUTPUT_SCHEMA: dict[str, Any] = {
    "type": "object",
    "properties": {
        "summary": {
            "type": "string",
            "description": "提案全体の短い要約（日本語、1〜3文）",
        },
        "insertions": {
            "type": "array",
            "description": "挿入する数式のリスト",
            "items": {
                "type": "object",
                "properties": {
                    "sheet": {"type": "string", "description": "対象シート名"},
                    "cell": {"type": "string", "description": "A1形式のセル位置（例: B12）"},
                    "formula": {"type": "string", "description": "= で始まる Excel 数式"},
                    "label": {"type": "string", "description": "この数式が何を計算するかの短い説明"},
                },
                "required": ["sheet", "cell", "formula", "label"],
                "additionalProperties": False,
            },
        },
    },
    "required": ["summary", "insertions"],
    "additionalProperties": False,
}


def _truncate(value: Any, limit: int = 60) -> str:
    s = repr(value)
    return s if len(s) <= limit else s[: limit - 3] + "..."


def describe_workbook(wb: Workbook, sample_rows: int = 5) -> str:
    """LLMに渡すためのワークブック構造の要約を組み立てる。"""
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"## シート: {ws.title}")
        parts.append(f"- 行数: {ws.max_row}, 列数: {ws.max_column}")
        if ws.max_row == 0 or ws.max_column == 0:
            parts.append("- (空のシート)")
            continue

        headers = [
            f"{get_column_letter(c)}={_truncate(ws.cell(row=1, column=c).value)}"
            for c in range(1, ws.max_column + 1)
        ]
        parts.append("- ヘッダ (1行目): " + ", ".join(headers))

        sample = min(sample_rows, max(0, ws.max_row - 1))
        if sample > 0:
            parts.append(f"- データサンプル (2行目から{sample}行):")
            for r in range(2, 2 + sample):
                row = [
                    f"{get_column_letter(c)}{r}={_truncate(ws.cell(row=r, column=c).value)}"
                    for c in range(1, ws.max_column + 1)
                ]
                parts.append("  " + ", ".join(row))
    return "\n".join(parts)


def generate_formula_plan(
    client: anthropic.Anthropic,
    workbook_description: str,
    instruction: str,
) -> tuple[dict[str, Any], Any]:
    """Claudeを呼び、数式挿入計画(dict)と usage を返す。"""
    instruction_text = instruction.strip() or "(指示なし — 実用的な集計関数を自由に提案してください)"
    user_message = (
        "# ワークブック構造\n"
        f"{workbook_description}\n\n"
        "# ユーザー指示\n"
        f"{instruction_text}\n\n"
        "上記の構造に対して、適切なExcel数式の挿入計画をJSONで返してください。"
    )

    response = client.messages.create(
        model=MODEL,
        max_tokens=16000,
        thinking={"type": "adaptive"},
        system=[
            {
                "type": "text",
                "text": SYSTEM_PROMPT,
                "cache_control": {"type": "ephemeral"},
            }
        ],
        output_config={
            "format": {
                "type": "json_schema",
                "schema": OUTPUT_SCHEMA,
            }
        },
        messages=[{"role": "user", "content": user_message}],
    )

    text = next(b.text for b in response.content if b.type == "text")
    return json.loads(text), response.usage


_CELL_REF_RE = re.compile(r"^[A-Za-z]{1,3}[1-9][0-9]*$")


def apply_insertions(wb: Workbook, insertions: list[dict[str, str]]) -> list[str]:
    """各insertionをワークブックに適用。警告のリストを返す。"""
    warnings: list[str] = []
    for ins in insertions:
        sheet = ins.get("sheet", "")
        cell = (ins.get("cell") or "").upper().strip()
        formula = (ins.get("formula") or "").strip()

        if sheet not in wb.sheetnames:
            warnings.append(f"シート '{sheet}' が存在しません: {cell} {formula} をスキップ")
            continue
        if not _CELL_REF_RE.match(cell):
            warnings.append(f"セル参照 '{cell}' が不正です: スキップ")
            continue
        if not formula:
            warnings.append(f"{sheet}!{cell}: 数式が空です。スキップ")
            continue
        if not formula.startswith("="):
            formula = "=" + formula

        ws = wb[sheet]
        existing = ws[cell].value
        if existing is not None:
            warnings.append(
                f"{sheet}!{cell} には既存値 {existing!r} がありました（上書きしました）"
            )
        ws[cell] = formula
    return warnings
